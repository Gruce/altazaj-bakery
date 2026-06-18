# -*- coding: utf-8 -*-
"""
Generate db/seed/menu.json from the FineDine Excel export (then `npm run db:seed`).

Usage:
    python3 scripts/generate-menu.py [path/to/export.xlsx]

Defaults to ./menu-source.xlsx in the repo root. Requires `openpyxl`
(`pip3 install openpyxl`). Re-run whenever the menu changes, then refine
any awkward English strings in the PHRASES / WORDS maps below.
"""
import json, re, os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "menu-source.xlsx")


def _load_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Menu"] if "Menu" in wb.sheetnames else wb.worksheets[0]
    out = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        cat, name, desc, price, img = (tuple(r) + (None,) * 5)[:5]
        if not name:
            continue
        out.append({
            "category": (cat or "").strip(),
            "name": (name or "").strip(),
            "desc": (desc or "").strip() if desc else "",
            "price": price,
            "img": (img or "").strip(),
        })
    return out


RAW = _load_xlsx(SRC)

# ---------------------------------------------------------------------------
# Category metadata: slug, English name, emoji
# ---------------------------------------------------------------------------
CATEGORIES = [
    ("المناقيش",            "manakish",       "Manakish",          "🫓"),
    ("سناك منقوشة",         "manoushe-snacks","Manoushe Snacks",   "🌯"),
    ("الساندويجات",         "sandwiches",     "Sandwiches",        "🥪"),
    ("كرواسون",             "croissants",     "Croissants",        "🥐"),
    ("مني كرواسون",         "mini-croissants","Mini Croissants",   "🥐"),
    ("فطائر",               "pastries",       "Pastries",          "🥧"),
    ("فطائر كبيرة",         "large-pastries", "Large Pastries",    "🥟"),
    ("دونات",               "donuts",         "Donuts",            "🍩"),
    ("الكيك",               "cakes",          "Cakes",             "🍰"),
    ("قوالب كيك",           "whole-cakes",    "Whole Cakes",       "🎂"),
    ("براونيز",             "brownies",       "Brownies",          "🍫"),
    ("سبونج كيك",           "sponge-cake",    "Sponge Cake",       "🍰"),
    ("مافن كيك",            "muffins",        "Muffins",           "🧁"),
    ("مافن كيك بالكريمة",   "cream-muffins",  "Cream Muffins",     "🧁"),
    ("تشيز كيك",            "cheesecake",     "Cheesecake",        "🍰"),
    ("الحلو العربي",        "arabic-sweets",  "Arabic Sweets",     "🍯"),
    ("البيتيفور",           "petit-four",     "Petit Four",        "🍪"),
    ("الكعك",               "kaak",           "Kaak",              "🥨"),
    ("المخبوزات",           "breads",         "Breads & Bakes",    "🍞"),
    ("مخبوزات صحية",        "healthy-bakes",  "Healthy Bakes",     "🌾"),
    ("مشروبات الساخنة",     "hot-drinks",     "Hot Drinks",        "☕"),
    ("موهيتو",              "mojito",         "Mojito",            "🍃"),
    ("مشروبات مثلجة",       "iced-drinks",    "Iced Drinks",       "🧊"),
    ("ميلك شيك",            "milkshakes",     "Milkshakes",        "🥤"),
    ("سموذي",               "smoothies",      "Smoothies",         "🥤"),
    ("فرابتشينو",           "frappuccino",    "Frappuccino",       "🥤"),
    ("العصائر الطبيعية",    "fresh-juices",   "Fresh Juices",      "🍊"),
    ("بيدا",                "bida",           "Bida",              "🫓"),
]
CAT_SLUG = {ar: slug for ar, slug, _, _ in CATEGORIES}
CAT_EN   = {ar: en   for ar, _, en, _ in CATEGORIES}
CAT_EMO  = {ar: emo  for ar, _, _, emo in CATEGORIES}

# ---------------------------------------------------------------------------
# Multi-word phrases (applied first, longest-first)
# ---------------------------------------------------------------------------
PHRASES = {
    "لحم بالعجين الطازج": "Fresh Lahm Bi Ajin",
    "لحم بالعجين": "Lahm Bi Ajin", "لحم بعجين": "Lahm Bi Ajin",
    "كت كات": "KitKat", "توت بري": "Cranberry", "راس بيري": "Raspberry",
    "بلو كوراسو": "Blue Curacao", "بلو هاواي": "Blue Hawaii",
    "فوريه نوار": "Black Forest", "فوريه بلانش": "White Forest",
    "سان سبستيان": "San Sebastian", "سبونج كيك": "Sponge Cake",
    "ميلك شيك": "Milkshake", "تشيز كيك": "Cheesecake", "كب كيك": "Cupcake",
    "بول كيك": "Cake Balls", "كريم بولز": "Cream Balls",
    "ايس كريم": "Ice Cream", "ايس كوفي": "Iced Coffee", "ايس كافيه": "Iced Coffee",
    "هوت دوغ": "Hot Dog", "هوت جوكليت": "Hot Chocolate", "هوت شوكلت": "Hot Chocolate",
    "فلات وايت": "Flat White", "سبانش لاتيه": "Spanish Latte",
    "قهوة تركية": "Turkish Coffee", "قهوة تركي": "Turkish Coffee",
    "قهوة عثمانية": "Ottoman Coffee", "ريد فيلفت": "Red Velvet",
    "سويس رول": "Swiss Roll", "كعب الغزال": "Kaab El Ghazal",
    "تراب الملوك": "Turab Al-Molouk", "جوز الهند": "Coconut",
    "ماء الزهر": "Blossom Water", "وماء الزهر": "& Blossom Water",
    "دبس فليفلة": "Pepper Molasses", "دبس الفليفلة": "Pepper Molasses",
    "باشن فروت": "Passion Fruit", "بلو بيري": "Blueberry",
    "حبة البركة": "Black Seed", "خبز فرنسي": "French Bread",
    "خبز بركر": "Burger Buns", "خبز مني بركر": "Mini Burger Buns",
    "خبز لوف": "Loaf Bread", "شوكلا شقراء": "Blonde Chocolate",
    "الشوكلا الشقراء": "Blonde Chocolate", "الشوكلا البيضاء": "White Chocolate",
    "شوكلا بيضاء": "White Chocolate", "جوكليت دارك": "Dark Chocolate",
    "سفن اب": "7Up", "ايس تي": "Iced Tea", "ايس شاي": "Iced Tea",
    "صب ابيض": "White Sob", "صب اسود": "Black Sob",
    "نعناع طبيعي": "Fresh Mint", "موهيتو الطازج": "Signature Mojito",
    "بطاطا مع اللحم": "Potato & Meat", "شرحات دجاج": "Chicken Slices",
    "صدر دجاج": "Chicken Breast", "لحم عراقي": "Iraqi Meat",
    "حبش مدخن": "Smoked Turkey", "لحم ببروني": "Pepperoni",
    "سناك منقوشة": "Manoushe Snack", "فطائر لحم": "Meat Pastry",
    "خبز توست": "Toast Bread",
    "متعدد الحبوب والدقيق الابيض": "Multigrain & White Flour",
    "متعدد الحبوب": "Multigrain", "دقيق النخالة": "Bran Flour",
    "بدقيق النخالة": "with Bran Flour", "الدقيق الابيض": "White Flour",
    "بالدقيق الابيض": "with White Flour", "بالدقيق الاسمر": "with Whole-Wheat Flour",
    "الدقيق الاسمر": "Whole-Wheat Flour", "دقيق ابيض": "White Flour",
}

# ---------------------------------------------------------------------------
# Single-token dictionary (normalized form -> English)
# ---------------------------------------------------------------------------
WORDS = {
    # connectors / quantity / units
    "مع":"with","و":"&","من":"from","بدون":"without","نصف":"Half","ربع":"Quarter",
    "غرام":"g","قطع":"pcs","قطعة":"pc","حبة":"pc","حبات":"pcs","كوب":"Cup","علبة":"Box",
    "باكيت":"Baguette","بطعم":"flavored","بنكهة":"flavored","اضافات":"Add-ons","مشكل":"Assorted","مكس":"Mix","بيدا":"Bida",
    "صغير":"Small","كبير":"Large","كبيرة":"Large","مني":"Mini","متعدد":"Multi","الخاصة":"Signature","خاص":"Special",
    # bases / bakery
    "خبز":"Bread","كيك":"Cake","سبونج":"Sponge","كرواسون":"Croissant","الكرواسون":"Croissant",
    "دونات":"Donut","الدونات":"Donut","فطيرة":"Pie","فطائر":"Pastry","منقوشة":"Manoushe","سناك":"Snack",
    "ساندويج":"Sandwich","توست":"Toast","بايغل":"Bagel","بريوش":"Brioche","رول":"Roll","لوف":"Loaf",
    "تارت":"Tart","التارت":"Tart","براونيز":"Brownies","قالب":"Whole Cake","كب":"Cup","مافن":"Muffin","المافن":"Muffin",
    "بسكويت":"Biscuit","بسوكيت":"Biscuit","كوكيز":"Cookies","سابليه":"Sable","سابيله":"Sable","سابيه":"Sable",
    "معمول":"Maamoul","غريبة":"Ghraybeh","كعك":"Kaak","جرك":"Jurk","فلفون":"Felfoun","بيتيفور":"Petit Four","بيتفور":"Petit Four",
    "بركر":"Burger","بيتزا":"Pizza","اكلير":"Eclair","كلير":"Eclair","ماكرون":"Macaron","دنش":"Danish",
    "برازق":"Barazek","البرازق":"Barazek","عثمانية":"Ottoman","بعلبكية":"Baalbaki","كرامبل":"Crumble","تيراميسو":"Tiramisu",
    "اوبرا":"Opera","فاج":"Fudge","فيلفت":"Velvet","نوار":"Noir","بلانش":"Blanche","فوريه":"Forest",
    "بول":"Ball","بولز":"Balls","كرات":"Balls","عجينة":"Dough","عجين":"Dough","عحين":"Dough","حشوة":"Filling","سادة":"Plain",
    "مغطس":"Dipped","مغطسة":"Dipped","محشو":"Stuffed","محشوة":"Stuffed","محلى":"Sweetened","صب":"Sob",
    # proteins / savory
    "لحم":"Meat","اللحم":"Meat","دجاج":"Chicken","صدر":"Breast","حبش":"Turkey","سجق":"Sujuk",
    "طاووق":"Tawook","ببروني":"Pepperoni","مارتديلا":"Mortadella","تونة":"Tuna","شاورما":"Shawarma","كفتة":"Kofta",
    "بيض":"Egg","عراقي":"Iraqi","مدخن":"Smoked","متبل":"Marinated","شرحات":"Slices","شرائح":"Slices","كلوب":"Club",
    "فاهيتا":"Fajita","فيلادلفيا":"Philadelphia","بيكاسو":"Picasso","كاساديا":"Quesadilla","مكسيكي":"Mexican","المكسيكي":"Mexican",
    # cheese / dairy
    "جبن":"Cheese","جبنة":"Cheese","الجبن":"Cheese","تشيز":"Cheese","شيدر":"Cheddar","الشيدر":"Cheddar",
    "حلوم":"Halloumi","لبنة":"Labneh","كريمة":"Cream","كريم":"Cream","الكريمة":"Cream","حليب":"Milk","حليبى":"Milky","لبن":"Yogurt",
    "كاستر":"Custard","كاسترد":"Custard","سمن":"Ghee","السمن":"Ghee","حيواني":"Animal","الحيواني":"Animal",
    # veg / herbs
    "طماطة":"Tomato","بصل":"Onion","خس":"Lettuce","خيار":"Cucumber","كبيس":"Pickles","ثوم":"Garlic",
    "معدنوس":"Parsley","فلفل":"Pepper","بارد":"Capsicum","فليفلة":"Pepper","فليلفة":"Pepper","جزر":"Carrot",
    "بطاطا":"Potato","ذرة":"Corn","زيتون":"Olive","سبانغ":"Spinach","سبانش":"Spinach","فطر":"Mushroom","خضار":"Vegetables",
    "اخضر":"Green","احمر":"Red","ريد":"Red","دارك":"Dark","سويس":"Swiss","اسباني":"Spanish",
    # zaatar / spices / nuts
    "زعتر":"Zaatar","سمسم":"Sesame","سماق":"Sumac","محمرة":"Muhammara","اوريغانو":"Oregano","زنجبيل":"Ginger",
    "سينامون":"Cinnamon","القرفة":"Cinnamon","دارسين":"Cinnamon","الهيل":"Cardamom","قرفة":"Cinnamon",
    "جوز":"Walnut","لوز":"Almond","بندق":"Hazelnut","فستق":"Pistachio","بستاشيو":"Pistachio","بستاسيو":"Pistachio",
    "زبيب":"Raisins","تمر":"Dates","عسل":"Honey","حلاوة":"Halawa",
    # sweet / chocolate
    "شوكلا":"Chocolate","شوكولا":"Chocolate","جوكليت":"Chocolate","كراميل":"Caramel","فانيلا":"Vanilla",
    "اوريو":"Oreo","لوتس":"Lotus","سنيكرز":"Snickers","توكس":"Twix","نوتلا":"Nutella","كلاكسي":"Galaxy",
    "ليون":"Lion","كرانش":"Crunch","سفن":"Seven","اب":"Up","كوراسو":"Curacao","سكر":"Sugar","سيروب":"Syrup","شقراء":"Blonde",
    # fruit
    "فراولة":"Strawberry","فروالة":"Strawberry","توت":"Berry","بيري":"Berry","فواكه":"Fruits","فاكهة":"Fruit",
    "ليمون":"Lemon","برتقال":"Orange","البرتقال":"Orange","برتفال":"Orange","تفاح":"Apple","ابل":"Apple",
    "اناناس":"Pineapple","مانجو":"Mango","موز":"Banana","كيوي":"Kiwi","خوخ":"Peach","فروت":"Fruit","باشن":"Passion","هاواي":"Hawaii",
    # drinks
    "اسبريسو":"Espresso","اسريسو":"Espresso","قهوة":"Coffee","كوفي":"Coffee","لاتيه":"Latte","موكا":"Mocha",
    "كابتشينو":"Cappuccino","كورتادو":"Cortado","امريكانو":"Americano","نسكافيه":"Nescafe","شاي":"Tea","تركية":"Turkish","تركي":"Turkish",
    "موهيتو":"Mojito","سموذي":"Smoothie","فرابشينو":"Frappuccino","فرابتشينو":"Frappuccino","عصير":"Juice",
    "نعناع":"Mint","ثلج":"Ice","ايس":"Iced","مثلج":"Iced","مثلجة":"Iced","ساخن":"Hot","بلو":"Blue","باودر":"Powder","مركز":"Concentrated",
    # bread / dough qualifiers
    "دقيق":"Flour","ابيض":"White","اسود":"Black","اسمر":"Whole Wheat","اشقر":"Blonde","نخالة":"Bran","حبوب":"Grains","الحبوب":"Grains",
    "نخاله":"Bran","فرنسي":"French","انكلش":"English","صحي":"Healthy","الصحي":"Healthy","طبيعي":"Fresh","خليط":"Blend","خليطنا":"Our Blend",
    "بيضاء":"White","سوداء":"Black","ملون":"Colorful","الملون":"Colorful","حبيبات":"Sprinkles","موالح":"Savory","حلو":"Sweet","عربي":"Arabic",
    # sauces
    "مايونيز":"Mayo","صوص":"Sauce","كاتشب":"Ketchup","دبس":"Molasses","رز":"Rice","ماء":"Water","الزهر":"Blossom",
    "هند":"Coconut","الهند":"Coconut","الملوك":"Kings","تراب":"Turab","كورتينا":"Cortina","بلو":"Blue","كرت":"Kit","كرامبل":"Crumble",
}

PREFIX_RULES = [
    ("وبال","& "), ("بال","with "), ("وبـ","& "), ("وال","& "), ("فال",""),
    ("بـ","with "), ("ال",""), ("و","& "), ("ب","with "),
]
SMALL = {"with","&","and","from","a","the","of","g","ml","pcs"}

def normalize(t):
    t = t.strip()
    t = re.sub("[إأآا]","ا",t)
    t = t.replace("ى","ي").replace("ؤ","و").replace("ئ","ي").replace("ـ","")
    t = re.sub("[ً-ْ]","",t)  # harakat
    return t

WORDS = {normalize(k): v for k, v in WORDS.items()}

def translit(t):
    m = {"ا":"a","ب":"b","ت":"t","ث":"th","ج":"j","ح":"h","خ":"kh","د":"d","ذ":"dh",
         "ر":"r","ز":"z","س":"s","ش":"sh","ص":"s","ض":"d","ط":"t","ظ":"z","ع":"a",
         "غ":"gh","ف":"f","ق":"q","ك":"k","ل":"l","م":"m","ن":"n","ه":"h","و":"w",
         "ي":"y","ة":"a","ء":"","ئ":"y","ؤ":"w"}
    return "".join(m.get(c,c) for c in t).capitalize()

def tok_en(tok):
    if not tok: return ""
    if re.fullmatch(r"[A-Za-z0-9&%()/+.-]+", tok): return tok      # already latin
    if tok.isdigit(): return tok
    n = normalize(tok)
    if n in WORDS: return WORDS[n]
    for pre, eng in PREFIX_RULES:
        if n.startswith(pre) and len(n) > len(pre)+1:
            root = n[len(pre):]
            if root in WORDS:
                return (eng + WORDS[root]).strip()
    return translit(n)

def smart_title(s):
    out=[]
    for i,w in enumerate(s.split()):
        lw=w.lower()
        if i>0 and lw in SMALL: out.append(lw)
        elif w and w[0].isalpha() and w[:1].islower(): out.append(w[:1].upper()+w[1:])
        else: out.append(w)
    return " ".join(out)

def apply_phrases(s):
    for ar in sorted(PHRASES, key=len, reverse=True):
        s = s.replace(ar, " "+PHRASES[ar]+" ")
    return s

def translate_field(s, title=True):
    if not s: return ""
    s = apply_phrases(s)
    segs = re.split(r"\s*[,،]\s*", s)
    out=[]
    for seg in segs:
        words=[tok_en(w) for w in seg.split()]
        seg_en=" ".join(w for w in words if w).strip()
        seg_en=re.sub(r"\s+"," ",seg_en)
        if title: seg_en=smart_title(seg_en)
        if seg_en: out.append(seg_en)
    return ", ".join(out)

# ---------------------------------------------------------------------------
# Featured signature picks (by ar-name substring)
# ---------------------------------------------------------------------------
FEATURED = [
    "لحم بالعجين الطازج","منقوشة زعتر مع الجبن","كرواسون شوكلا","تشيز كيك فراولة",
    "سان سبستيان","براونيز","دونات سادة","موهيتو الطازج كبير","سبانش لاتيه",
    "ريد فيلفت","كت كات","سموذي مانجو","معمول تمر","فرابشينو كراميل",
]

def is_featured(name):
    return any(f in name for f in FEATURED)

# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------
items=[]
cat_count={}
seen_ids=set()
for idx, it in enumerate(RAW):
    cat=it["category"]; name=it["name"]; desc=it["desc"]; price=it["price"]; img=it["img"]
    m=re.search(r"/items/(\d+)/", img)
    iid=("i"+m.group(1)) if m else f"x{idx}"
    if iid in seen_ids: iid=f"{iid}-{idx}"
    seen_ids.add(iid)
    cat_count[cat]=cat_count.get(cat,0)+1
    items.append({
        "id": iid,
        "cat": CAT_SLUG.get(cat, "other"),
        "ar": {"n": name, "d": desc},
        "en": {"n": translate_field(name, True), "d": translate_field(desc, True)},
        "price": price,
        "img": img,
        "feat": is_featured(name),
    })

categories=[{"id":slug,"ar":ar,"en":en,"emoji":emo,"count":cat_count.get(ar,0)}
            for ar,slug,en,emo in CATEGORIES]

out={
    "brand":{
        "ar":{"name":"الطازج بيكري","tag":"طازج كل يوم","city":"البصرة، شارع الجزائر"},
        "en":{"name":"Altazaj Bakery","tag":"Freshly baked, every day","city":"Basra · Algeria St."},
        "currency":{"ar":"د.ع","en":"IQD"},
        "facebook":"https://www.facebook.com/altazajbakery/",
    },
    "categories":categories,
    "items":items,
}

dst = os.path.join(ROOT, "db", "seed", "menu.json")
os.makedirs(os.path.dirname(dst), exist_ok=True)
json.dump(out, open(dst, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print(f"Wrote {dst} — run `npm run db:seed` to load it into SQLite.")

# Coverage report
ar_left=set()
for it in items:
    for f in (it["en"]["n"], it["en"]["d"]):
        for w in f.split():
            if re.search(r"[؀-ۿ]", w): ar_left.add(w)
print("Items:",len(items),"Categories:",len(categories),"Featured:",sum(1 for i in items if i['feat']))
print("Untranslated AR tokens remaining:", len(ar_left))
print(sorted(ar_left)[:40])
print("\nSAMPLE TRANSLATIONS:")
for s in [0,15,60,120,180,250,300,330]:
    if s<len(items):
        it=items[s]
        print(f"  AR {it['ar']['n']} :: {it['ar']['d']}")
        print(f"  EN {it['en']['n']} :: {it['en']['d']}\n")
print("Wrote", dst, os.path.getsize(dst), "bytes")
